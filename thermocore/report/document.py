"""
Geração do relatório final de inspeção, em formato DOCX.

Esta é a última etapa do pipeline descrito na metodologia do projeto
(... → detecção → visualização → relatório). O relatório é montado
inteiramente a partir de um `AnalysisRecord` já preenchido (tudo que aparece no relatório é rastreável até os dados de
entrada) e,
opcionalmente, de uma imagem anotada já gerada por
`thermocore.visualization.annotate.save_annotated_image`.

Por que DOCX e não PDF diretamente: `python-docx` gera o arquivo DOCX
de forma inteiramente nativa em Python, sem depender de nenhum
programa externo instalado na máquina do usuário (Word, LibreOffice
etc.), o que é essencial para portabilidade — este pacote roda tanto
em Windows quanto em Linux/macOS, e não podemos presumir que um
conversor de PDF esteja disponível. Quem precisar de PDF pode abrir o
`.docx` gerado no Word ou no LibreOffice e exportar para PDF
manualmente (Arquivo → Salvar como → PDF) — isso é responsabilidade da
etapa de "publicação", separada da geração de dados deste projeto.
Geração de PDF nativa (por exemplo, via biblioteca dedicada) é
trabalho futuro, listado no README.

Este módulo NUNCA sobrescreve a imagem original nem o arquivo
anotado — apenas os lê (quando fornecidos) para montar um documento
novo.
"""

from __future__ import annotations

from pathlib import Path

from thermocore.report.provenance import AnalysisRecord


def _format_value(value: object) -> str:
    """Formata um valor do dicionário de parâmetros/resultados para exibição."""
    if isinstance(value, float):
        return f"{value:.2f}"
    if isinstance(value, bool):
        return "Sim" if value else "Não"
    return str(value)


def _fmt_pct_range(value: object) -> str:
    return f"{value[0]:.1f}% a {value[1]:.1f}%" if isinstance(value, (list, tuple)) and len(value) == 2 else str(value)


_HIDDEN_RESULT_KEYS = {"avisos", "sensibilidade_referencia"}


def generate_inspection_report_docx(
    record: AnalysisRecord,
    output_path: str | Path,
    annotated_image_path: str | Path | None = None,
) -> Path:
    """
    Monta um relatório de inspeção em DOCX a partir de `record` e,
    opcionalmente, embute `annotated_image_path` (a imagem já marcada
    com o ponto quente/referência, gerada por `save_annotated_image`).

    Parâmetros
    ----------
    record: registro de rastreabilidade já preenchido com os dados
        desta análise (ver `thermocore.report.provenance.AnalysisRecord`).
    output_path: caminho do arquivo .docx a ser criado. O diretório pai
        é criado automaticamente se não existir.
    annotated_image_path: caminho de uma imagem já gerada (JPG/PNG) a
        ser embutida no relatório. Se None ou se o arquivo não existir,
        o relatório é gerado sem imagem, e isso é indicado no texto —
        nunca falha silenciosamente nem insere uma imagem em branco.

    Retorna o `Path` do arquivo `.docx` criado.
    """
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()

    doc.add_heading("Relatório de Inspeção Termográfica", level=1)
    subtitle = doc.add_paragraph()
    subtitle.add_run(
        f"Gerado automaticamente por thermocore v{record.software_version}"
    ).italic = True

    # --- Aviso de destaque quando o dado é sintético ---------------------
    if record.is_synthetic:
        warning = doc.add_paragraph()
        run = warning.add_run(
            "⚠ ESTE RELATÓRIO FOI GERADO A PARTIR DE DADOS SINTÉTICOS. "
            "NÃO REPRESENTA UMA INSPEÇÃO REAL E NÃO DEVE SER USADO PARA "
            "NENHUMA DECISÃO OPERACIONAL."
        )
        run.bold = True
        run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)

    # --- Identificação da análise -----------------------------------------
    doc.add_heading("Identificação da análise", level=2)
    id_table = doc.add_table(rows=0, cols=2)
    id_table.style = "Light Grid Accent 1"
    id_rows = [
        ("Arquivo de origem", record.source_path),
        ("Dado sintético", "Sim" if record.is_synthetic else "Não"),
        ("Data/hora da análise (UTC)", record.analyzed_at),
        ("Método", record.method),
        ("Versão do software (thermocore)", record.software_version),
    ]
    for label, value in id_rows:
        row_cells = id_table.add_row().cells
        row_cells[0].text = label
        row_cells[1].text = str(value)
        row_cells[0].paragraphs[0].runs[0].bold = True

    # --- Parâmetros utilizados --------------------------------------------
    if record.parameters:
        doc.add_heading("Parâmetros utilizados", level=2)
        param_table = doc.add_table(rows=0, cols=2)
        param_table.style = "Light Grid Accent 1"
        for key, value in record.parameters.items():
            row_cells = param_table.add_row().cells
            row_cells[0].text = key
            row_cells[1].text = _format_value(value)
            row_cells[0].paragraphs[0].runs[0].bold = True

    # --- Resultados ---------------------------------------------------------
    doc.add_heading("Resultados", level=2)
    if record.results:
        result_table = doc.add_table(rows=0, cols=2)
        result_table.style = "Light Grid Accent 1"
        for key, value in record.results.items():
            if key in _HIDDEN_RESULT_KEYS:
                continue  # exibidos em seções próprias, abaixo
            row_cells = result_table.add_row().cells
            row_cells[0].text = key
            if key == "eta_faixa_pct":
                row_cells[1].text = _fmt_pct_range(value)
            elif isinstance(value, (list, tuple)):
                row_cells[1].text = " a ".join(str(v) for v in value)
            else:
                row_cells[1].text = _format_value(value)
            row_cells[0].paragraphs[0].runs[0].bold = True
    else:
        doc.add_paragraph("Nenhum resultado registrado neste AnalysisRecord.")

    # --- Sensibilidade à referência ---------------------------------------------
    sensitivity = record.results.get("sensibilidade_referencia") if record.results else None
    if sensitivity:
        doc.add_heading("Sensibilidade à escolha da referência", level=2)
        doc.add_paragraph(
            "η e a classe de severidade recalculados para diferentes percentis de "
            "temperatura do material usados como referência (T2)."
        )
        sens_table = doc.add_table(rows=1, cols=5)
        sens_table.style = "Light Grid Accent 1"
        for i, head in enumerate(["Percentil", "T referência (°C)", "ΔT (°C)", "η (%)", "Severidade"]):
            sens_table.rows[0].cells[i].text = head
        for row in sensitivity:
            cells = sens_table.add_row().cells
            cells[0].text = f"P{row['percentil']:.0f}"
            cells[1].text = f"{row['t_referencia_c']:.1f}"
            cells[2].text = f"{row['delta_t_c']:.1f}"
            cells[3].text = f"{row['eta_pct']:.1f}"
            cells[4].text = str(row["severidade"])

    # --- Avisos desta análise ----------------------------------------------------
    analysis_warnings = record.results.get("avisos") if record.results else None
    if analysis_warnings:
        doc.add_heading("Avisos desta análise", level=2)
        for item in analysis_warnings:
            warn_par = doc.add_paragraph(style="List Bullet")
            run = warn_par.add_run(item)
            run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)

    # --- Imagem anotada -----------------------------------------------------
    doc.add_heading("Imagem anotada", level=2)
    image_path = Path(annotated_image_path) if annotated_image_path else None
    if image_path is not None and image_path.exists():
        doc.add_picture(str(image_path), width=Inches(6))
    else:
        note = doc.add_paragraph()
        note.add_run(
            "Nenhuma imagem anotada foi fornecida ou encontrada para esta "
            "análise."
        ).italic = True

    # --- Avisos e limitações -------------------------------------------------
    doc.add_heading("Avisos e limitações", level=2)
    limitations = [
        "O ponto de referência (T2) é escolhido automaticamente como o pixel do "
        "material estimado cuja temperatura é a mais próxima de um percentil baixo "
        "(padrão P10) — não é uma região escolhida por termografista, e η e a "
        "classe de severidade dependem dessa escolha (ver a tabela de sensibilidade). "
        "Confirme a referência em componente ou fase homóloga antes de qualquer "
        "decisão operacional.",
        "A região de material é estimada por limiar de Otsu e tende a subsegmentar "
        "partes do equipamento mais frias que seu núcleo mais quente; ela assume "
        "fundo distinto do equipamento.",
        "Esta versão do sistema marca apenas um ponto quente e um ponto de "
        "referência por imagem — não trata múltiplos pontos quentes no "
        "mesmo equipamento nem distingue reflexo especular de anomalia "
        "térmica real.",
        "Os parâmetros de medição (emissividade, distância, T0) são os configurados "
        "na câmera no momento da captura; T0 não é uma medição independente do "
        "ambiente.",
        "A classificação de severidade é determinística, baseada em limiares "
        "normativos (DL/T 664-2016, categoria padrão \"outros equipamentos de "
        "corrente\") e em um ΔT mínimo configurável; é uma triagem preliminar e "
        "não substitui o laudo de um termografista. Não envolve nenhum modelo de "
        "aprendizado de máquina.",
        "A \"severidade\" é um critério térmico, e não uma matriz de risco: a ABNT NBR "
        "15866 recomenda que a prioridade de ação considere também segurança, perda "
        "de produção, criticidade e tendência do equipamento, o que esta ferramenta "
        "não faz. A norma não prescreve limiares numéricos.",
        "Não há correção de carga nem de vento, e a emissividade é única por imagem "
        "(a configurada na câmera); metais de baixa emissividade, comuns em conexões, "
        "dificultam a avaliação quantitativa.",
    ]
    for item in limitations:
        doc.add_paragraph(item, style="List Bullet")

    footer_paragraph = doc.add_paragraph()
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer_paragraph.add_run(
        "Relatório gerado automaticamente por thermocore — "
        "projeto de TCC de Igor Pantoja, Engenharia de Computação, UFPA."
    )
    footer_run.font.size = Pt(8)
    footer_run.italic = True

    doc.save(str(output_path))
    return output_path
