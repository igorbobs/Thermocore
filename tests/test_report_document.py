from docx import Document

from thermocore.report.document import generate_inspection_report_docx
from thermocore.report.provenance import AnalysisRecord


def _make_record(is_synthetic: bool = True) -> AnalysisRecord:
    return AnalysisRecord(
        source_path="synthetic-hotspot",
        is_synthetic=is_synthetic,
        method="gaussian_denoise+threshold_segmentation+eta_DLT664",
        parameters={"threshold_celsius": 60.0, "gaussian_sigma": 1.0},
        results={"t_max_celsius": 93.2, "eta_pct": 97.8, "severity": "emergencial"},
    )


def test_generate_report_creates_docx_file(tmp_path):
    record = _make_record()
    output_path = tmp_path / "relatorio.docx"

    result_path = generate_inspection_report_docx(record, output_path)

    assert result_path == output_path
    assert output_path.exists()

    # Confere que é um .docx válido e legível.
    doc = Document(str(output_path))
    assert len(doc.paragraphs) > 0


def test_generate_report_flags_synthetic_data_prominently(tmp_path):
    record = _make_record(is_synthetic=True)
    output_path = tmp_path / "relatorio_sintetico.docx"
    generate_inspection_report_docx(record, output_path)

    doc = Document(str(output_path))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "SINTÉTICOS" in full_text or "SINTÉTICO" in full_text


def test_generate_report_without_synthetic_flag_for_real_data(tmp_path):
    record = _make_record(is_synthetic=False)
    output_path = tmp_path / "relatorio_real.docx"
    generate_inspection_report_docx(record, output_path)

    doc = Document(str(output_path))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "NÃO REPRESENTA UMA INSPEÇÃO REAL" not in full_text


def test_generate_report_includes_parameters_and_results(tmp_path):
    record = _make_record()
    output_path = tmp_path / "relatorio_dados.docx"
    generate_inspection_report_docx(record, output_path)

    doc = Document(str(output_path))
    all_cells_text = []
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                all_cells_text.append(cell.text)
    combined = "\n".join(all_cells_text)

    assert "threshold_celsius" in combined
    assert "eta_pct" in combined
    assert "emergencial" in combined


def test_generate_report_handles_missing_image_gracefully(tmp_path):
    record = _make_record()
    output_path = tmp_path / "relatorio_sem_imagem.docx"

    # Caminho de imagem que não existe — não deve levantar exceção.
    result_path = generate_inspection_report_docx(
        record, output_path, annotated_image_path=tmp_path / "nao_existe.jpg"
    )
    assert result_path.exists()


def test_generate_report_embeds_real_annotated_image(tmp_path):
    from thermocore.io.synthetic import make_hotspot_thermogram
    from thermocore.analysis.stats import locate_extreme
    from thermocore.visualization.annotate import save_annotated_image

    t = make_hotspot_thermogram()
    hotspot = locate_extreme(t.celsius, mode="max")
    reference = locate_extreme(t.celsius, mode="min")
    image_path = tmp_path / "anotada.jpg"
    save_annotated_image(
        t.celsius, hotspot=hotspot, reference=reference,
        output_path=image_path, delta_t_celsius=hotspot.value_celsius - reference.value_celsius,
    )

    record = _make_record()
    output_path = tmp_path / "relatorio_com_imagem.docx"
    generate_inspection_report_docx(record, output_path, annotated_image_path=image_path)

    doc = Document(str(output_path))
    # Uma imagem embutida gera uma relação (InlineShape) no documento.
    assert len(doc.inline_shapes) == 1
