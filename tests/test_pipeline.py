"""Pipeline ponta a ponta (arquivo → análise → registro → saídas) e regressão com imagens reais."""

from pathlib import Path

import pytest

from thermocore.analysis import AnomalySeverity
from thermocore.exceptions import ThermocoreError
from thermocore.pipeline import analyze_file, export_outputs

FIX = Path(__file__).parent / "fixtures"

pytestmark = pytest.mark.skipif(
    not (FIX / "IR_2497_t640_sample.jpg").exists(), reason="Fixtures reais não encontradas."
)

# Valores de referência (v0.2.0, FLIR T640). Servem de regressão e de verificação
# entre plataformas: rode `pytest` no Windows e no Linux e compare.
GOLDEN = {
    "IR_2497_t640_sample.jpg": dict(t1=59.81, t2=39.76, eta=77.66, sev=AnomalySeverity.GENERAL, hot=(288, 352)),
    "IR_2501_t640_lines.jpg": dict(t1=49.77, t2=36.12, eta=86.52, sev=AnomalySeverity.MAJOR, hot=(250, 184)),
    "IR_2503_t640_panel.jpg": dict(t1=36.62, t2=35.50, eta=42.64, sev=AnomalySeverity.NORMAL, hot=(4, 639)),
    "IR_2511_t640_switchgear.jpg": dict(t1=53.30, t2=37.37, eta=82.48, sev=AnomalySeverity.MAJOR, hot=(147, 289)),
}


@pytest.mark.parametrize("name", list(GOLDEN))
def test_real_images_reproduce_reference_values(name):
    g = GOLDEN[name]
    r = analyze_file(FIX / name).result
    assert r.hotspot.value_celsius == pytest.approx(g["t1"], abs=0.02)
    assert (r.hotspot.row, r.hotspot.col) == g["hot"]
    assert r.reference.value_celsius == pytest.approx(g["t2"], abs=0.02)
    assert r.eta_pct == pytest.approx(g["eta"], abs=0.05)
    assert r.severity == g["sev"]


def test_validation_image_matches_camera_legend_within_rounding():
    t = analyze_file(FIX / "IR_2497_t640_sample.jpg").thermogram
    assert float(t.celsius.max()) == pytest.approx(59.6, abs=0.3)


def test_record_has_relative_source_and_all_new_fields():
    fa = analyze_file(FIX / "IR_2497_t640_sample.jpg", source_label="tests/fixtures/IR_2497_t640_sample.jpg")
    rec = fa.record
    assert rec.source_path == "tests/fixtures/IR_2497_t640_sample.jpg"
    assert rec.software_version == "0.2.0"
    for key in ("eta_faixa_pct", "severity_faixa", "sensibilidade_referencia", "avisos", "separabilidade_otsu"):
        assert key in rec.results
    assert len(rec.results["sensibilidade_referencia"]) == 5
    assert rec.parameters["erosion_iterations"] == 4


def test_default_source_label_is_file_name_not_absolute_path():
    fa = analyze_file(FIX / "IR_2503_t640_panel.jpg")
    assert fa.record.source_path == "IR_2503_t640_panel.jpg"


def test_unsupported_file_raises_thermocore_error(tmp_path):
    bad = tmp_path / "x.txt"
    bad.write_text("não é imagem")
    with pytest.raises(ThermocoreError):
        analyze_file(bad)


def test_export_outputs_creates_new_files_and_leaves_original_untouched(tmp_path):
    src = FIX / "IR_2501_t640_lines.jpg"
    before = src.read_bytes()
    fa = analyze_file(src)
    image, report = export_outputs(fa, tmp_path / "a.jpg", tmp_path / "r.docx")
    assert image.exists() and report.exists()
    assert src.read_bytes() == before


def test_report_lists_warnings_sensitivity_and_current_limitations(tmp_path):
    from docx import Document

    fa = analyze_file(FIX / "IR_2501_t640_lines.jpg")
    _, report = export_outputs(fa, tmp_path / "a.jpg", tmp_path / "r.docx")
    doc = Document(str(report))
    text = "\n".join(p.text for p in doc.paragraphs) + "\n".join(
        c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert "Avisos desta análise" in text
    assert "Sensibilidade à escolha da referência" in text
    assert "fora da região de material" in text                     # aviso do pico de 66 °C
    assert "pixel mais frio de toda a imagem" not in text           # limitação antiga removida


def test_report_states_class_is_thermal_criterion_not_risk_matrix(tmp_path):
    from docx import Document

    fa = analyze_file(FIX / "IR_2497_t640_sample.jpg", max_admissible_celsius=50.0)
    assert fa.record.results["mta_excedida"] is True
    _, report = export_outputs(fa, tmp_path / "a.jpg", tmp_path / "r.docx")
    text = "\n".join(p.text for p in Document(str(report)).paragraphs)
    assert "matriz de risco" in text and "NBR 15866" in text
    assert "correção de carga nem de vento" in text
